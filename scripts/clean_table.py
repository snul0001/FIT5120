"""
Clean AI-impact table.

Turns the detailed output of extract_core into one tidy table per task:
ANZSCO code, Occupation, Task core, AI impact group, AI impact level,
Human need type, Short explanation.

Run extract_core.py first (it writes output/bespoketable_core_meaning_mapped.xlsx).
Run:  python scripts/clean_table.py
Needs: pandas, openpyxl
"""

# ----------------------------------------------------------------------------
# 1. Paths and settings
# ----------------------------------------------------------------------------
from pathlib import Path
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent   # folder this script is in
BASE_DIR = SCRIPT_DIR if (SCRIPT_DIR / "output").exists() else SCRIPT_DIR.parent
OUTPUT_DIR = BASE_DIR / "output"

# input: the workbook written by extract_core (mapped version if it exists)
SOURCE = OUTPUT_DIR / "bespoketable_core_meaning_mapped.xlsx"
if not SOURCE.exists():
    SOURCE = OUTPUT_DIR / "bespoketable_core_meaning.xlsx"
CLEAN_XLSX = OUTPUT_DIR / "ai_impact_clean.xlsx"
CLEAN_CSV = OUTPUT_DIR / "ai_impact_clean.csv"

# ---- settings you can change -------------------------------------------------
AUTOMATION_CUT = 0.6    # automation score >= this            -> AI-automated
AUGMENT_CUT = 0.6       # else augmentation score >= this     -> AI-assisted   (else Human-led)
HIGH_CUT = 0.7          # within a group, deciding score >= this -> High, otherwise Medium
INCLUDE_SCORES = False  # True adds the two raw scores to the clean table

comb = pd.read_excel(SOURCE, sheet_name="Combined")
print("read", len(comb), "tasks,", comb["ANZSCO unit code"].nunique(), "occupations from", SOURCE.name)

# ----------------------------------------------------------------------------
# 2. Group, level and explanation
# ----------------------------------------------------------------------------
HUMAN_ROLE = {   # plain-language reading of each human-need type
    "Physical / hands-on": "the hands-on work stays with people",
    "Interpersonal": "the people side stays human",
    "Judgment / expertise": "human judgment and expertise are still needed",
    "Oversight / final check": "a person checks the result",
}
GENERIC_HUMAN = {"human", "oversight", "intervention", "hands", "on", "physical", "manual", "labor", "labour",
                 "involvement", "input", "judgment", "expertise", "skilled", "presence", "interaction", "element",
                 "touch", "designed", "capable", "support", "skills", "skill", "nuanced", "complex"}


def top(cell, n=3, drop_generic=False):
    """First n phrases of a '; '-separated cell as a comma list ('' if empty)."""
    if not isinstance(cell, str) or not cell.strip():
        return ""
    ps = [p for p in cell.split("; ") if p]
    if drop_generic:      # "human oversight" adds nothing once the role sentence has said it
        ps = [p for p in ps if not set(p.replace("-", " ").split()) <= GENERIC_HUMAN]
    return ", ".join(ps[:n])


def classify(row):
    """Group + level from the two scores; human need type and explanation from the extracted columns."""
    A, G = row["Automation task score"], row["Augmentation task score"]
    ai_auto = top(row["Automation: AI can (core)"])
    ai_what = top(row["Augmentation: AI action (what)"])
    ai_how = top(row["Augmentation: AI method (how)"], 2)
    how = f" (by {ai_how})" if ai_how else ""

    if A >= AUTOMATION_CUT:
        group, level = "AI-automated", ("High" if A >= HIGH_CUT else "Medium")
        htype = row["Automation: Human need type"]
        hum = top(row["Automation: Human still needed (core)"], 2, True)
        text = f"AI can do most of this itself: {ai_auto or ai_what}."
    elif G >= AUGMENT_CUT:
        group, level = "AI-assisted", ("High" if G >= HIGH_CUT else "Medium")
        htype = row["Augmentation: Human need type"]
        hum = top(row["Augmentation: Human need (detail)"], 2, True)
        text = f"You do this with AI support: AI helps with {ai_what}{how}."
    else:
        group, level = "Human-led", "Low"
        htype = row["Automation: Human need type"]
        hum = top(row["Automation: Human still needed (core)"], 3, True)
        text = "Mostly human work." + (f" AI's help is limited to {ai_what}." if ai_what else "")

    role = HUMAN_ROLE.get(htype)
    if role:              # oversight is boilerplate, so its detail is left out
        detail = hum if htype != "Oversight / final check" else ""
        text += f" {role[0].upper()}{role[1:]}" + (f": {detail}." if detail else ".")
    return pd.Series({"AI impact group": group, "AI impact level": level,
                      "Human need type": htype, "Short explanation": text})

# ----------------------------------------------------------------------------
# 3. Build the clean table
# ----------------------------------------------------------------------------
both = comb.dropna(subset=["Automation task score", "Augmentation task score"]).copy()
if len(both) < len(comb):
    print(len(comb) - len(both), "tasks skipped: only one score available")

labels = both.apply(classify, axis=1)
clean = pd.DataFrame({
    "ANZSCO code": both["ANZSCO unit code"].astype(int),
    "Occupation": both["ANZSCO unit title"],
    "Task core": both["Task core"],
    "AI impact group": labels["AI impact group"],
    "AI impact level": labels["AI impact level"],
    "Human need type": labels["Human need type"],
    "Short explanation": labels["Short explanation"],
})
if INCLUDE_SCORES:
    clean.insert(3, "Automation score", both["Automation task score"])
    clean.insert(4, "Augmentation score", both["Augmentation task score"])

# sort: occupation, then most AI-affected tasks first
g_order = {"AI-automated": 0, "AI-assisted": 1, "Human-led": 2}
l_order = {"High": 0, "Medium": 1, "Low": 2}
clean = (clean.assign(_g=clean["AI impact group"].map(g_order), _l=clean["AI impact level"].map(l_order))
              .sort_values(["ANZSCO code", "_g", "_l"], kind="stable").drop(columns=["_g", "_l"]).reset_index(drop=True))

print(clean["AI impact group"].value_counts().to_string())

# ----------------------------------------------------------------------------
# 4. How many tasks in each group and level
# ----------------------------------------------------------------------------
summary = (clean.groupby(["AI impact group", "AI impact level"]).size().rename("Tasks").reset_index())
summary["Share"] = (summary["Tasks"] / len(clean) * 100).round(1).astype(str) + "%"
summary = summary.sort_values(["AI impact group", "AI impact level"],
                              key=lambda s: s.map({**g_order, **l_order})).reset_index(drop=True)
print("\nTasks per group and level:")
print(summary.to_string(index=False))

# ----------------------------------------------------------------------------
# 5. Profile per occupation
# ----------------------------------------------------------------------------
# share of each occupation's tasks at each group + level
by_occ = (pd.crosstab([clean["ANZSCO code"], clean["Occupation"]],
                      clean["AI impact group"] + " - " + clean["AI impact level"], normalize="index") * 100).round(0)
cols = [c for c in ["AI-automated - High", "AI-automated - Medium", "AI-assisted - High",
                    "AI-assisted - Medium", "Human-led - Low"] if c in by_occ.columns]
by_occ = by_occ[cols].add_prefix("% ").reset_index()
by_occ.insert(2, "Tasks", clean.groupby("ANZSCO code").size().values)

# ----------------------------------------------------------------------------
# 6. Save (Excel and CSV)
# ----------------------------------------------------------------------------
rules = pd.DataFrame([
    ("AI-automated", "High", f"automation score >= {HIGH_CUT}"),
    ("AI-automated", "Medium", f"automation score {AUTOMATION_CUT} to below {HIGH_CUT}"),
    ("AI-assisted", "High", f"automation below {AUTOMATION_CUT} and augmentation score >= {HIGH_CUT}"),
    ("AI-assisted", "Medium", f"automation below {AUTOMATION_CUT} and augmentation {AUGMENT_CUT} to below {HIGH_CUT}"),
    ("Human-led", "Low", f"automation below {AUTOMATION_CUT} and augmentation below {AUGMENT_CUT}"),
], columns=["AI impact group", "AI impact level", "Rule"]).merge(summary, how="left").fillna({"Tasks": 0, "Share": "0%"})

WIDTHS = {"ANZSCO code": 9, "Occupation": 34, "Task core": 46, "AI impact group": 15, "AI impact level": 11,
          "Human need type": 22, "Short explanation": 80, "Rule": 60}
FILL = {"High": "F8CBAD", "Medium": "FFE699", "Low": "C6E0B4"}

with pd.ExcelWriter(CLEAN_XLSX, engine="openpyxl") as xw:
    for name, d in [("AI_impact", clean), ("Rules_and_counts", rules), ("By_occupation", by_occ)]:
        d.to_excel(xw, sheet_name=name, index=False)
        ws = xw.sheets[name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(d.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 16)
            h = ws.cell(1, i)
            h.font = Font(bold=True, color="FFFFFF", name="Arial")
            h.fill = PatternFill("solid", fgColor="1F4E78")
            h.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        lvl = list(d.columns).index("AI impact level") + 1 if "AI impact level" in d.columns else None
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.font = Font(name="Arial", size=10)
            if lvl:           # colour the level cell: High / Medium / Low
                cell = row[lvl - 1]
                cell.fill = PatternFill("solid", fgColor=FILL.get(cell.value, "FFFFFF"))

clean.to_csv(CLEAN_CSV, index=False)
print("saved", CLEAN_XLSX)
print("saved", CLEAN_CSV)
